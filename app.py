import os
import re
import sqlite3
from datetime import datetime
from functools import wraps
from pathlib import Path

from flask import Flask, flash, redirect, render_template, request, send_file, session, url_for
from werkzeug.security import check_password_hash, generate_password_hash
from openpyxl import Workbook, load_workbook

BASE_DIR = Path(__file__).resolve().parent
app = Flask(__name__)
app.config.update(
    SECRET_KEY=os.environ.get("SECRET_KEY", "change-this-secret-before-deployment"),
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
)
DATABASE = os.environ.get("DATABASE_PATH", str(BASE_DIR / "campuscare.db"))
LOGIN_EXCEL = BASE_DIR / "Campus360_Login_Data.xlsx"
ADMIN_EMAIL = os.environ.get("ADMIN_EMAIL", "admin@campus.360").strip().lower()
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "Campus360@Admin2026")
STATUSES = ["Submitted", "Assigned", "In Progress", "Resolved", "Closed"]
CATEGORIES = ["Hostel / Room", "Campus", "Classroom Facilities", "Electrical", "Cleaning", "Water", "Wi-Fi", "Other"]
DEFAULT_WORKERS = [
    ("Ravi Kumar", "Electrical", "9876543210"),
    ("Suresh Reddy", "Plumbing", "9876543211"),
    ("Arjun Kumar", "Cleaning", "9876543212"),
    ("Naveen Kumar", "Maintenance", "9876543213"),
    ("Mahesh", "Wi-Fi & Network", "9876543214"),
]


def valid_campus_email(email):
    return bool(re.fullmatch(r"[^@\s]+@campus\.360", (email or "").strip().lower()))


def db():
    con = sqlite3.connect(DATABASE, timeout=15)
    con.row_factory = sqlite3.Row
    return con


def init_db():
    with db() as con:
        con.execute("""CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL,
            email TEXT UNIQUE NOT NULL, password TEXT NOT NULL,
            role TEXT NOT NULL DEFAULT 'student', mobile TEXT DEFAULT '',
            student_id TEXT DEFAULT '', department TEXT DEFAULT '', year TEXT DEFAULT '', room TEXT DEFAULT '')""")
        con.execute("""CREATE TABLE IF NOT EXISTS workers (
            id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL,
            department TEXT NOT NULL, mobile TEXT NOT NULL,
            user_id INTEGER UNIQUE)""")
        # Upgrade databases created by older versions without losing existing records.
        worker_columns = {row[1] for row in con.execute("PRAGMA table_info(workers)").fetchall()}
        if "user_id" not in worker_columns:
            con.execute("ALTER TABLE workers ADD COLUMN user_id INTEGER")
        con.execute("""CREATE TABLE IF NOT EXISTS complaints (
            id INTEGER PRIMARY KEY AUTOINCREMENT, complaint_id TEXT UNIQUE NOT NULL,
            user_id INTEGER NOT NULL, title TEXT NOT NULL, category TEXT NOT NULL,
            location TEXT DEFAULT '', description TEXT DEFAULT '',
            status TEXT NOT NULL DEFAULT 'Submitted', worker_id INTEGER,
            created_at TEXT NOT NULL, updated_at TEXT NOT NULL)""")
        if not con.execute("SELECT id FROM users WHERE email=?", (ADMIN_EMAIL,)).fetchone():
            con.execute("INSERT INTO users(name,email,password,role) VALUES(?,?,?,?)",
                        ("Campus Administrator", ADMIN_EMAIL, generate_password_hash(ADMIN_PASSWORD), "admin"))
        if con.execute("SELECT COUNT(*) FROM workers").fetchone()[0] == 0:
            con.executemany("INSERT INTO workers(name,department,mobile) VALUES(?,?,?)", DEFAULT_WORKERS)


def log_login(name, email, role, status):
    headers = ["Date/Time", "Name", "Campus Email", "Role", "Login Status"]
    try:
        if LOGIN_EXCEL.exists():
            workbook = load_workbook(LOGIN_EXCEL)
            sheet = workbook.active
            if sheet.max_row == 1 and sheet.cell(1, 1).value is None:
                sheet.append(headers)
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Login Data"
            sheet.append(headers)
        sheet.append([datetime.now().strftime("%Y-%m-%d %H:%M:%S"), name, email, role, status])
        workbook.save(LOGIN_EXCEL)
    except (OSError, PermissionError):
        app.logger.exception("Could not write login audit workbook")


def current_user():
    return session.get("user")


def require(role=None):
    user = current_user()
    return bool(user and (role is None or user.get("role") == role))


def role_required(role=None):
    def decorator(view):
        @wraps(view)
        def wrapped(*args, **kwargs):
            if not require(role):
                flash("Please sign in with an authorized account.", "error")
                return redirect(url_for("home"))
            return view(*args, **kwargs)
        return wrapped
    return decorator


@app.after_request
def security_headers(response):
    response.headers.setdefault("X-Content-Type-Options", "nosniff")
    response.headers.setdefault("X-Frame-Options", "SAMEORIGIN")
    response.headers.setdefault("Referrer-Policy", "strict-origin-when-cross-origin")
    return response


@app.route("/")
def home():
    user = current_user()
    if user:
        if user.get("role") == "worker":
            return redirect(url_for("worker_portal"))
        return redirect(url_for("dashboard"))
    return render_template("index.html")


@app.post("/login")
def login():
    email = request.form.get("email", "").strip().lower()
    password = request.form.get("password", "")
    if not valid_campus_email(email):
        flash("Use your official @campus.360 account.", "error")
        return redirect(url_for("home"))
    with db() as con:
        user = con.execute("SELECT * FROM users WHERE email=?", (email,)).fetchone()
    if user and check_password_hash(user["password"], password):
        session.clear()
        session["user"] = {key: user[key] for key in user.keys() if key != "password"}
        log_login(user["name"], email, user["role"], "SUCCESS")
        flash("Welcome back. You are signed in.", "success")
        if user["role"] == "worker":
            return redirect(url_for("worker_portal"))
        return redirect(url_for("dashboard"))
    log_login(user["name"] if user else "Unknown", email, user["role"] if user else "unknown", "FAILED")
    flash("Invalid email or password.", "error")
    return redirect(url_for("home"))


@app.get("/logout")
def logout():
    session.clear()
    flash("You have been logged out.", "success")
    return redirect(url_for("home"))


@app.route("/dashboard")
@role_required()
def dashboard():
    user = current_user()
    if user["role"] == "worker":
        return redirect(url_for("worker_portal"))
    with db() as con:
        if user["role"] == "admin":
            complaints = con.execute("""SELECT c.*, u.name AS student, w.name AS worker
                FROM complaints c JOIN users u ON u.id=c.user_id
                LEFT JOIN workers w ON w.id=c.worker_id ORDER BY c.id DESC""").fetchall()
            students = con.execute("SELECT * FROM users WHERE role='student' ORDER BY id DESC").fetchall()
            workers = con.execute("SELECT * FROM workers ORDER BY id").fetchall()
            return render_template("admin.html", user=user, complaints=complaints,
                                   students=students, workers=workers, statuses=STATUSES)
        complaints = con.execute("""SELECT c.*, w.name AS worker FROM complaints c
            LEFT JOIN workers w ON w.id=c.worker_id WHERE c.user_id=? ORDER BY c.id DESC""",
                                 (user["id"],)).fetchall()
    return render_template("student.html", user=user, complaints=complaints, categories=CATEGORIES)


@app.post("/admin/create-student")
@role_required("admin")
def create_student():
    form = request.form
    name = form.get("name", "").strip()
    email = form.get("email", "").strip().lower()
    mobile = re.sub(r"\D", "", form.get("mobile", ""))
    password = form.get("password", "")
    student_id = form.get("student_id", "").strip()
    if not name or not valid_campus_email(email) or len(mobile) != 10 or len(password) < 8:
        flash("Enter a name, valid @campus.360 email, 10-digit mobile and password (at least 8 characters).", "error")
        return redirect(url_for("dashboard"))
    try:
        with db() as con:
            con.execute("""INSERT INTO users(name,email,password,role,mobile,student_id,department,year,room)
                VALUES(?,?,?,?,?,?,?,?,?)""", (name, email, generate_password_hash(password), "student", mobile,
                student_id, form.get("department", "").strip(), form.get("year", "").strip(), form.get("room", "").strip()))
        flash("Student account created successfully.", "success")
    except sqlite3.IntegrityError:
        flash("That email is already registered. Please use a different email.", "error")
    return redirect(url_for("dashboard"))


@app.post("/admin/create-worker")
@role_required("admin")
def create_worker():
    form = request.form
    name = form.get("name", "").strip()
    email = form.get("email", "").strip().lower()
    department = form.get("department", "").strip()
    mobile = re.sub(r"\D", "", form.get("mobile", ""))
    password = form.get("password", "")
    if not name or not valid_campus_email(email) or not department or len(mobile) != 10 or len(password) < 8:
        flash("Enter worker name, valid @campus.360 email, department, 10-digit mobile and password (at least 8 characters).", "error")
        return redirect(url_for("dashboard"))
    try:
        with db() as con:
            cur = con.execute(
                "INSERT INTO users(name,email,password,role,mobile) VALUES(?,?,?,?,?)",
                (name, email, generate_password_hash(password), "worker", mobile),
            )
            con.execute(
                "INSERT INTO workers(name,department,mobile,user_id) VALUES(?,?,?,?)",
                (name, department, mobile, cur.lastrowid),
            )
        flash("Worker account created. The worker can now sign in with the email and password.", "success")
    except sqlite3.IntegrityError:
        flash("That email is already registered. Use a different campus email.", "error")
    return redirect(url_for("dashboard"))


@app.get("/admin/login-data")
@role_required("admin")
def login_data():
    if not LOGIN_EXCEL.exists():
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Login Data"
        sheet.append(["Date/Time", "Name", "Campus Email", "Role", "Login Status"])
        workbook.save(LOGIN_EXCEL)
    return send_file(LOGIN_EXCEL, as_attachment=True, download_name="Campus360_Login_Data.xlsx")


@app.post("/admin/assign/<int:complaint_id>")
@role_required("admin")
def assign_worker(complaint_id):
    worker_raw = request.form.get("worker_id", "").strip()
    worker_id = int(worker_raw) if worker_raw.isdigit() else None
    with db() as con:
        if worker_id and not con.execute("SELECT id FROM workers WHERE id=?", (worker_id,)).fetchone():
            flash("Please choose a valid worker.", "error")
            return redirect(url_for("dashboard"))
        con.execute("UPDATE complaints SET worker_id=?, status=?, updated_at=? WHERE id=?",
                    (worker_id, "Assigned" if worker_id else "Submitted", datetime.now().strftime("%Y-%m-%d %H:%M:%S"), complaint_id))
    flash("Worker assignment updated.", "success")
    return redirect(url_for("dashboard"))


@app.post("/admin/status/<int:complaint_id>")
@role_required("admin")
def update_status(complaint_id):
    status = request.form.get("status", "Submitted")
    if status not in STATUSES:
        flash("Invalid complaint status.", "error")
        return redirect(url_for("dashboard"))
    with db() as con:
        con.execute("UPDATE complaints SET status=?, updated_at=? WHERE id=?",
                    (status, datetime.now().strftime("%Y-%m-%d %H:%M:%S"), complaint_id))
    flash("Complaint status updated.", "success")
    return redirect(url_for("dashboard"))


@app.post("/complaint")
@role_required("student")
def complaint():
    form = request.form
    title = form.get("title", "").strip()
    category = form.get("category", "Other")
    location = form.get("location", "").strip()
    description = form.get("description", "").strip()
    if not title or not description or category not in CATEGORIES:
        flash("Enter a title and description and choose a valid category.", "error")
        return redirect(url_for("dashboard"))
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    complaint_id = "CC-" + datetime.now().strftime("%y%m%d%H%M%S%f")[-16:]
    with db() as con:
        con.execute("""INSERT INTO complaints(complaint_id,user_id,title,category,location,description,created_at,updated_at)
            VALUES(?,?,?,?,?,?,?,?)""", (complaint_id, current_user()["id"], title, category, location, description, now, now))
    flash(f"Complaint {complaint_id} submitted successfully.", "success")
    return redirect(url_for("dashboard"))


@app.get("/worker")
@role_required("worker")
def worker_portal():
    user = current_user()
    with db() as con:
        worker = con.execute("SELECT * FROM workers WHERE user_id=?", (user["id"],)).fetchone()
        if not worker:
            flash("Your worker profile is not linked yet. Contact the campus administrator.", "error")
            return render_template("worker.html", user=user, jobs=[])
        jobs = con.execute("""SELECT c.*, u.name AS student FROM complaints c
            JOIN users u ON u.id=c.user_id WHERE c.worker_id=? ORDER BY c.id DESC""",
            (worker["id"],)).fetchall()
    return render_template("worker.html", user=user, jobs=jobs)


@app.post("/worker/status/<int:complaint_id>")
@role_required("worker")
def worker_status(complaint_id):
    status = request.form.get("status", "In Progress")
    if status not in ["In Progress", "Resolved"]:
        flash("Invalid status.", "error")
        return redirect(url_for("worker_portal"))
    user = current_user()
    with db() as con:
        worker = con.execute("SELECT id FROM workers WHERE user_id=?", (user["id"],)).fetchone()
        if not worker:
            flash("Worker profile not found. Contact the administrator.", "error")
            return redirect(url_for("worker_portal"))
        cur = con.execute("UPDATE complaints SET status=?, updated_at=? WHERE id=? AND worker_id=?",
            (status, datetime.now().strftime("%Y-%m-%d %H:%M:%S"), complaint_id, worker["id"]))
        if cur.rowcount == 0:
            flash("That complaint is not assigned to your account.", "error")
            return redirect(url_for("worker_portal"))
    flash("Work status updated successfully.", "success")
    return redirect(url_for("worker_portal"))


@app.get("/health")
def health():
    return {"status": "ok", "app": "Campus 360"}, 200


init_db()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "5000")), debug=os.environ.get("FLASK_DEBUG") == "1")
